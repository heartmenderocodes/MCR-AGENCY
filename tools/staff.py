"""
Staff, ajustes internos e información personal.

Fuentes internas que lee (cualquier .xlsx de la carpeta --interno):
  · hoja "Staff"              → personas del equipo: cargo, estado, cumpleaños, ingreso
  · hoja "Ajustes_Creadores"  → decisiones sobre creadores (creator_id, campo, valor, motivo)
Las hojas de información personal del maestro llevan celdas de captura (amarillas) y fórmulas bloqueadas.
"""
import datetime as dt
import os
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STAFF_HEADERS = {
    "NOMBRE": "nombre", "CARGO": "cargo", "ESTADO": "estado", "ÁREA": "area", "AREA": "area",
    "CUMPLE DÍA": "cumple_dia", "CUMPLE DIA": "cumple_dia", "CUMPLE MES": "cumple_mes",
    "AÑO NACIMIENTO": "nac_anio", "INGRESO AÑO": "ing_anio", "INGRESO MES": "ing_mes", "INGRESO DÍA": "ing_dia",
    "INGRESO DIA": "ing_dia", "FECHA SALIDA": "salida", "NOTAS": "notas", "ALIAS": "alias",
}
STAFF_TEMPLATE_COLS = ["NOMBRE", "CARGO", "ESTADO", "ÁREA", "CUMPLE DÍA", "CUMPLE MES", "AÑO NACIMIENTO",
                       "INGRESO AÑO", "INGRESO MES", "INGRESO DÍA", "FECHA SALIDA", "NOTAS"]
ESTADOS = ("Activo", "Ex-equipo", "Por confirmar")
PERSONAL_FIELDS = ("cumple_dia", "cumple_mes", "nac_anio", "ing_anio", "ing_mes", "ing_dia")
INPUT_FILL = PatternFill("solid", fgColor="FFF8D6")
FORMULA_FILL = PatternFill("solid", fgColor="F3ECFF")
FORMULA_FONT = Font(name="Arial", size=10, color="5500CC")


def _header_map(ws, mapping, marker):
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        norm = [str(v).strip().upper() if v else "" for v in vals]
        if marker in norm:
            return r, {mapping[v]: i for i, v in enumerate(norm, 1) if v in mapping}
    return None, {}


def _int(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return None
    try:
        return int(float(str(v).strip()))
    except ValueError:
        return None


def load_internal_adjustments(paths):
    """Devuelve (staff_sources, creator_adjustments). Cada staff_source = (path, filas)."""
    staff, adjustments = [], []
    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            title = ws.title.strip().lower()
            if title in ("staff", "equipo", "personal"):
                hr, cols = _header_map(ws, STAFF_HEADERS, "NOMBRE")
                if not hr:
                    continue
                rows = []
                for r in range(hr + 1, ws.max_row + 1):
                    name = ws.cell(r, cols["nombre"]).value
                    if not name or not str(name).strip():
                        continue
                    row = {k: ws.cell(r, c).value for k, c in cols.items()}
                    row["nombre"] = str(name).strip()
                    for f in PERSONAL_FIELDS:
                        if f in row:
                            row[f] = _int(row[f])
                    rows.append(row)
                staff.append((path, rows))
            elif title.startswith("ajustes"):
                hr, cols = _header_map(ws, {"CREATOR_ID": "creator_id", "CAMPO": "campo", "VALOR": "valor",
                                            "MOTIVO": "motivo"}, "CREATOR_ID")
                if not hr:
                    continue
                for r in range(hr + 1, ws.max_row + 1):
                    cid = ws.cell(r, cols["creator_id"]).value
                    if cid:
                        adjustments.append(dict(creator_id=str(cid).strip(), campo=ws.cell(r, cols["campo"]).value,
                                                valor=ws.cell(r, cols["valor"]).value,
                                                motivo=ws.cell(r, cols.get("motivo", 0)).value if cols.get("motivo") else None,
                                                fuente=os.path.basename(path)))
    return staff, adjustments


def load_previous_master(path):
    """Lee lo capturado a mano en las celdas de entrada de un maestro anterior."""
    if not path or not os.path.exists(path):
        return {}, {}
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for sheet in ("06_Personal_Staff", "07_Personal_Creadores"):
        data = {}
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            hdr = {str(c.value): i for i, c in enumerate(ws[4], 1) if c.value}
            label_to_field = {lab: f for f, lab, kind in STAFF_PERSONAL_COLS + CREATOR_PERSONAL_COLS if kind == "input"}
            for r in range(5, ws.max_row + 1):
                pid = ws.cell(r, 1).value
                if not pid:
                    continue
                vals = {label_to_field[lab]: _int(ws.cell(r, i).value) for lab, i in hdr.items() if lab in label_to_field}
                data[str(pid)] = {k: v for k, v in vals.items() if v is not None}
        out.append(data)
    return out[0], out[1]


def key(s):
    import unicodedata
    s = "".join(c for c in unicodedata.normalize("NFD", str(s or "")) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s.lower())


def merge_staff(team, staff_sources, previous, src, resolver, role_area):
    """Combina el equipo de la web (nivel 2) con las hojas Staff internas (nivel 1)."""
    by_key = {key(t["nombre"]): t for t in team}
    first = {}
    for t in team:
        first.setdefault(key(str(t["nombre"]).split()[0]), []).append(t)
    cands = {}  # person_id -> field -> [(valor, fuente)]

    def bucket(t):
        return cands.setdefault(t["person_id"], {})

    for t in team:
        b = bucket(t)
        for f in ("cargo", "estado", "area"):
            b.setdefault(f, []).append((t.get(f), t["fuente"]))
    for sid, rows in staff_sources:
        for r in rows:
            t = by_key.get(key(r["nombre"])) or (by_key.get(key(r["alias"])) if r.get("alias") else None)
            if t and r.get("alias") and key(t["nombre"]) == key(r["alias"]):
                t["nombre"] = r["nombre"]  # el nombre interno reemplaza al alias de la web
                by_key[key(r["nombre"])] = t
            if not t:
                fk = first.get(key(r["nombre"].split()[0]), [])
                t = fk[0] if len(fk) == 1 and len(r["nombre"].split()) == 1 else None
            if not t:
                t = dict(person_id=f"P{len(team) + 1:03d}", nombre=r["nombre"], cargo=None, area=None, nivel_org="Equipo",
                         estado=None, foto=None, fuente=sid, notas="Alta desde archivo interno (no está en la web)")
                team.append(t)
                by_key[key(t["nombre"])] = t
            b = bucket(t)
            for f in ("cargo", "estado", "area", "salida", "notas") + PERSONAL_FIELDS:
                if r.get(f) not in (None, ""):
                    b.setdefault(f, []).append((r[f], sid))
    for pid, vals in previous.items():
        t = next((x for x in team if x["person_id"] == pid), None)
        if t:
            for f, v in vals.items():
                bucket(t).setdefault(f, []).append((v, "S40"))
    for t in team:
        b = cands.get(t["person_id"], {})
        for f in ("cargo", "estado", "area", "salida") + PERSONAL_FIELDS:
            v, s = resolver.pick("Staff", t["person_id"], f, b.get(f, []), numeric=f in PERSONAL_FIELDS)
            t[f] = v
            t[f"fuente_{f}"] = s
        notas_int = [v for v, s in b.get("notas", []) if s != "S04"]
        if notas_int:
            t["notas"] = "; ".join(filter(None, [t.get("notas")] + notas_int))
        t["estado"] = normalize_estado(t.get("estado"))
        t["area"] = t.get("area") or role_area(t.get("cargo"))
    return team


def normalize_estado(v):
    s = str(v or "").lower()
    if not s:
        return "Por confirmar"
    if "ex" in s or "baja" in s or "salió" in s:
        return "Ex-equipo"
    if "activo" in s:
        return "Activo"
    return "Por confirmar"


# ─────────────────────── hojas de información personal ───────────────────────
# (campo, etiqueta, tipo) · tipo: "data" (lo pone el generador), "input" (editable), "formula"

STAFF_PERSONAL_COLS = [
    ("person_id", "ID", "data"), ("nombre", "Nombre", "data"), ("cargo", "Cargo", "data"),
    ("cumple_dia", "Día", "input"), ("cumple_mes", "Mes", "input"), ("nac_anio", "Año nacimiento", "input"),
    ("cumple", "Cumpleaños", "formula"), ("edad", "Edad", "formula"), ("proximo_cumple", "Próximo cumpleaños", "formula"),
    ("dias_cumple", "Días para cumple", "formula"),
    ("ing_anio", "Ingreso año", "input"), ("ing_mes", "Ingreso mes", "input"), ("ing_dia", "Ingreso día", "input"),
    ("fecha_ingreso", "Fecha ingreso", "formula"), ("precision", "Precisión ingreso", "formula"),
    ("antiguedad", "Antigüedad", "formula"), ("proximo_aniversario", "Próximo aniversario", "formula"),
    ("dias_aniversario", "Días para aniversario", "formula"), ("notas", "Notas", "input"),
]

CREATOR_PERSONAL_COLS = [
    ("person_id", "ID", "data"), ("nombre", "Creador", "data"), ("nombre_real", "Nombre real", "data"),
    ("tier", "Tier", "data"), ("pais", "País", "data"),
    ("cumple_dia", "Día", "input"), ("cumple_mes", "Mes", "input"), ("nac_anio", "Año nacimiento", "input"),
    ("cumple", "Cumpleaños", "formula"), ("edad", "Edad", "formula"), ("proximo_cumple", "Próximo cumpleaños", "formula"),
    ("dias_cumple", "Días para cumple", "formula"),
    ("ing_anio", "Firma año", "input"), ("ing_mes", "Firma mes", "input"), ("ing_dia", "Firma día", "input"),
    ("fecha_ingreso", "Fecha firma", "formula"), ("precision", "Precisión firma", "formula"),
    ("antiguedad", "Antigüedad en MCR", "formula"), ("proximo_aniversario", "Próximo aniversario", "formula"),
    ("dias_aniversario", "Días para aniversario", "formula"), ("notas", "Notas", "input"),
]


def personal_formulas(columns, n):
    L = {f: get_column_letter(i) for i, (f, _, _) in enumerate(columns, 1)}
    d, m, y = f"{L['cumple_dia']}{n}", f"{L['cumple_mes']}{n}", f"{L['nac_anio']}{n}"
    ia, im, idd = f"{L['ing_anio']}{n}", f"{L['ing_mes']}{n}", f"{L['ing_dia']}{n}"
    fi, pc, pa = f"{L['fecha_ingreso']}{n}", f"{L['proximo_cumple']}{n}", f"{L['proximo_aniversario']}{n}"
    day_i = f'IF({idd}="",1,{idd})'
    return {
        "cumple": f'=IF(OR({d}="",{m}=""),"",TEXT(DATE(2000,{m},{d}),"dd/mm"))',
        "edad": f'=IF(OR({y}="",{d}="",{m}=""),"",DATEDIF(DATE({y},{m},{d}),TODAY(),"y"))',
        "proximo_cumple": f'=IF(OR({d}="",{m}=""),"",DATE(YEAR(TODAY())+IF(DATE(YEAR(TODAY()),{m},{d})<TODAY(),1,0),{m},{d}))',
        "dias_cumple": f'=IF({pc}="","",{pc}-TODAY())',
        "fecha_ingreso": f'=IF({ia}="","",DATE({ia},IF({im}="",1,{im}),{day_i}))',
        "precision": f'=IF({ia}="","",IF({im}="","Sólo año",IF({idd}="","Año y mes","Fecha exacta")))',
        "antiguedad": (f'=IF({fi}="","",IF({im}="",(YEAR(TODAY())-{ia})&" años (aprox.)",'
                       f'DATEDIF({fi},TODAY(),"y")&" a "&DATEDIF({fi},TODAY(),"ym")&" m"))'),
        "proximo_aniversario": (f'=IF(OR({fi}="",{im}=""),"",DATE(YEAR(TODAY())+'
                                f'IF(DATE(YEAR(TODAY()),{im},{day_i})<TODAY(),1,0),{im},{day_i}))'),
        "dias_aniversario": f'=IF({pa}="","",{pa}-TODAY())',
    }


def staff_personal_rows(team):
    rows = []
    for t in team:
        fc = t.get("fuente_cumple_dia")
        fi = t.get("fuente_ing_anio")
        rows.append(dict(person_id=t["person_id"], nombre=t["nombre"], cargo=t.get("cargo"),
                         cumple_dia=t.get("cumple_dia"), cumple_mes=t.get("cumple_mes"), nac_anio=t.get("nac_anio"),
                         ing_anio=t.get("ing_anio"), ing_mes=t.get("ing_mes"), ing_dia=t.get("ing_dia"),
                         fuente_cumple=fc, fuente_ingreso=fi, notas=None))
    return sorted(rows, key=lambda r: r["person_id"])


def creator_personal_rows(creators, previous):
    rows = []
    for c in creators:
        prev = previous.get(c["creator_id"], {})
        rows.append(dict(person_id=c["creator_id"], nombre=c["nombre_publico"], nombre_real=c["nombre_real"],
                         tier=c["tier"], pais=c["pais"],
                         cumple_dia=c["cumple_dia"] or prev.get("cumple_dia"), cumple_mes=c["cumple_mes"] or prev.get("cumple_mes"),
                         nac_anio=c.get("nac_anio") or prev.get("nac_anio"), ing_anio=c.get("firma_anio") or prev.get("ing_anio"),
                         ing_mes=c.get("firma_mes") or prev.get("ing_mes"), ing_dia=c.get("firma_dia") or prev.get("ing_dia"),
                         fuente_cumple="S02" if c["cumple_dia"] else ("S40" if prev.get("cumple_dia") else None),
                         fuente_ingreso="S40" if prev.get("ing_anio") else None,
                         notas=None))
    return rows


def finish_personal_sheet(ws, columns, n_rows, hdr=4):
    """Rellena fórmulas, colorea entradas, valida y protege la hoja (sin contraseña)."""
    first, last = hdr + 1, hdr + max(n_rows, 1)
    L = {f: i for i, (f, _, _) in enumerate(columns, 1)}
    for n in range(first, last + 1):
        for f, formula in personal_formulas(columns, n).items():
            c = ws.cell(n, L[f], formula)
            c.fill = FORMULA_FILL
            c.font = FORMULA_FONT
            if f in ("proximo_cumple", "proximo_aniversario", "fecha_ingreso"):
                c.number_format = "DD/MM/YYYY"
        for f, _, kind in columns:
            cell = ws.cell(n, L[f])
            if kind == "input":
                cell.fill = INPUT_FILL
                cell.protection = Protection(locked=False)
    for f, label, kind in columns:
        hc = ws.cell(hdr, L[f])
        if kind == "input":
            hc.fill = PatternFill("solid", fgColor="8833FF")
        elif kind == "formula":
            hc.fill = PatternFill("solid", fgColor="0F0F0F")
    rules = [("cumple_dia", 1, 31), ("cumple_mes", 1, 12), ("nac_anio", 1940, 2015), ("ing_anio", 2015, 2035),
             ("ing_mes", 1, 12), ("ing_dia", 1, 31)]
    for f, lo, hi in rules:
        col = get_column_letter(L[f])
        dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor fuera de rango", error=f"Escribe un número entre {lo} y {hi}.")
        dv.add(f"{col}{first}:{col}{last + 200}")
        ws.add_data_validation(dv)
    for n in range(last + 1, last + 201):  # filas extra editables para altas
        for f, _, kind in columns:
            ws.cell(n, L[f]).protection = Protection(locked=False) if kind != "formula" else Protection(locked=True)
    protect(ws)


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.autoFilter = False
    p.sort = False
    p.formatColumns = False
    p.formatRows = False
    p.selectLockedCells = False
    p.selectUnlockedCells = False
