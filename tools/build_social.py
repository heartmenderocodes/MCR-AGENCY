#!/usr/bin/env python3
"""
Biblia de social de MCR: marca, línea editorial, calendario mensual y solicitudes de diseño.

Uso:
  python3 build_social.py --maestro <MCR_MAESTRO recalculado.xlsx> --inicio 2026-11-01 --meses 6 --out <salida.xlsx>

La información de la empresa se toma del maestro (debe venir recalculado para leer los valores de las fórmulas).
"""
import argparse
import calendar
import datetime as dt

import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import layout
from layout import BLACK, FONT, GREEN, MUTED, PURPLE, PURPLE_200, PURPLE_400, PURPLE_50, PURPLE_600, WHITE

MESES = "Enero Febrero Marzo Abril Mayo Junio Julio Agosto Septiembre Octubre Noviembre Diciembre".split()
DIAS = "Lun Mar Mié Jue Vie Sáb Dom".split()
SLOTS = 4

PLATAFORMAS = ["Instagram", "LinkedIn", "Instagram + LinkedIn", "TikTok"]
FORMATOS = ["Historia", "Historia larga", "Reel", "Post", "Carrusel", "Artículo LinkedIn", "Documento LinkedIn"]
TIPOS = ["Repost", "Trabajo con marcas", "Activaciones en vivo", "Eventos de creadores", "Premios a la agencia",
         "Artículos de Sergi", "Pláticas de Sergi", "Otro"]
ESTADOS = ["Por definir", "En diseño", "En revisión", "Aprobado", "Programado", "Publicado", "Cancelado"]
APROBACION = ["Pendiente", "Aprobado", "Cambios solicitados"]
SOLICITANTES = ["Sergi Cerrato", "Jocelyn Santiago", "Dulce Valentín"]
LIDER_DISENO = "Sergio Barranco"
DISENO = ["Sergio Barranco", "Elena", "Diego Galaviz", "Diseñador 3 (por contratar)"]
RESPONSABLES = ["Dulce Valentín", "Jocelyn Santiago", "Sergi Cerrato", "Sergio Barranco", "Kevin Moss"]
PRIORIDAD = ["Baja", "Media", "Alta", "Urgente"]
STATUS_DISENO = ["En espera", "En proceso", "En revisión", "Cambios", "Aprobado", "Terminada", "Bloqueo"]
MEDIDAS = ["1080×1920 · 9:16 (historia / reel)", "1080×1350 · 4:5 (post / carrusel)", "1080×1080 · 1:1",
           "1200×627 · LinkedIn imagen", "1080×1350 · LinkedIn documento (PDF)", "1920×1080 · 16:9"]

THIN = Side(style="thin", color="E4DCF5")
DAY_TOP = Side(style="medium", color=PURPLE_200)


# ─────────────────────────── lectura del maestro ───────────────────────────

def read_table(wb, name):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    hdr = [c.value for c in ws[4]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=5) if r[0].value not in (None, "")]


def company(maestro):
    wb = openpyxl.load_workbook(maestro, data_only=True)
    sheet = {n.split("_", 1)[1]: n for n in wb.sheetnames if "_" in n}
    get = lambda n: read_table(wb, sheet.get(n, ""))
    return dict(agencia={r["Campo"]: r["Valor"] for r in get("Agencia")}, cifras=get("Cifras"), servicios=get("Servicios"),
                oficinas=get("Oficinas"), equipo=get("Equipo"), prensa=get("Prensa_Premios"),
                portafolio=get("Portafolio"), clientes=get("Clientes"))


# ─────────────────────────── utilidades de hoja ───────────────────────────

def cell(ws, r, c, v=None, bold=False, color=BLACK, size=10, fill=None, wrap=False, align=None, italic=False):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    x.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return x


def section(ws, r, title, ncols):
    for c in range(1, ncols + 1):
        ws.cell(r, c).fill = PatternFill("solid", fgColor=BLACK)
    cell(ws, r, 1, title.upper(), bold=True, color=WHITE, fill=BLACK)
    ws.row_dimensions[r].height = 22
    return r + 1


def mini_table(ws, r, headers, rows, wrap_cols=()):
    for i, h in enumerate(headers, 1):
        cell(ws, r, i, h, bold=True, color=WHITE, fill=PURPLE)
    r += 1
    for k, row in enumerate(rows):
        for i, v in enumerate(row, 1):
            x = cell(ws, r, i, v, wrap=(i - 1) in wrap_cols, fill=PURPLE_50 if k % 2 else None)
            x.border = Border(bottom=THIN)
        r += 1
    return r + 1


def dv_list(ws, rng, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.add(rng)
    ws.add_data_validation(dv)


# ─────────────────────────── hojas ───────────────────────────

def sheet_marca(wb, co):
    ws = wb.create_sheet("02_Marca")
    layout.banner(ws, "Marca", 5)
    ag = co["agencia"]
    r = 5
    r = section(ws, r, "Identidad", 5)
    ident = [("Nombre", ag.get("Nombre comercial")), ("Tagline", ag.get("Tagline")),
             ("Tagline secundario", ag.get("Tagline secundario")), ("Posicionamiento", ag.get("Posicionamiento")),
             ("Web", "mcr-agency.com"), ("Instagram", ag.get("Instagram")), ("LinkedIn", ag.get("LinkedIn")),
             ("Correo", ag.get("Correo general")), ("Oficinas", ag.get("Oficinas")), ("Idiomas", ag.get("Idiomas"))]
    r = mini_table(ws, r, ["Campo", "Valor"], [(k, v) for k, v in ident if v], wrap_cols=(1,))
    r = section(ws, r, "Paleta y tipografía", 5)
    cell(ws, r, 1, "Color", bold=True, color=WHITE, fill=PURPLE)
    cell(ws, r, 2, "Hex", bold=True, color=WHITE, fill=PURPLE)
    cell(ws, r, 3, "Uso", bold=True, color=WHITE, fill=PURPLE)
    cell(ws, r, 4, "Muestra", bold=True, color=WHITE, fill=PURPLE)
    r += 1
    palette = [("Morado MCR", PURPLE, "Color de marca, único"), ("Morado 600", PURPLE_600, "Hover, contraste"),
               ("Morado 400", PURPLE_400, "Highlights"), ("Morado 200", PURPLE_200, "Tintes sobre foto, glows"),
               ("Negro", BLACK, "Fondo base (nunca #000 puro)"), ("Superficie", "1E1E1E", "Cards y paneles"),
               ("Verde acento", GREEN, "Uso raro, sólo editorial"), ("Blanco", WHITE, "Texto principal sobre negro")]
    for name, hx, use in palette:
        cell(ws, r, 1, name)
        cell(ws, r, 2, f"#{hx}")
        cell(ws, r, 3, use)
        ws.cell(r, 4).fill = PatternFill("solid", fgColor=hx)
        ws.cell(r, 4).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        r += 1
    r += 1
    r = mini_table(ws, r, ["Tipografía", "Uso"], [("Tusker Grotesk", "Titulares y display"),
                                                   ("Helvetica Neue", "Texto, subtítulos y copy")])
    r = section(ws, r, "Cifras para comunicar", 5)
    r = mini_table(ws, r, ["Métrica", "Valor"], [(c["Métrica"], c["Valor"]) for c in co["cifras"]])
    r = section(ws, r, "Servicios", 5)
    r = mini_table(ws, r, ["Servicio", "Descripción"], [(s["Servicio"], s["Descripción"]) for s in co["servicios"]],
                   wrap_cols=(1,))
    exclusivas = [c["Cliente"] for c in co["clientes"] if c.get("Agencia exclusiva") == "Sí"]
    r = section(ws, r, "Clientes y casos", 5)
    r = mini_table(ws, r, ["Agencia exclusiva de", ""], [(", ".join(exclusivas), None)])
    r = mini_table(ws, r, ["Caso", "Cliente", "KPIs"],
                   [(p["Caso"], p["Cliente"], p["KPIs"]) for p in co["portafolio"]], wrap_cols=(2,))
    premios = [p for p in co["prensa"] if p.get("Tipo") != "Prensa"]
    r = section(ws, r, "Premios y reconocimientos", 5)
    r = mini_table(ws, r, ["Reconocimiento", "Organización", "Fecha"],
                   [(p["Título"], p.get("Medio / organización"), p.get("Fecha")) for p in premios], wrap_cols=(0,))
    equipo = {e["Nombre"]: e for e in co["equipo"]}
    voces = [("Sergi Cerrato", "Vocero · artículos y pláticas"), ("Jocelyn Santiago", "PR · aprueba todo lo que cruza a LinkedIn"),
             ("Dulce Valentín", "Community Manager · operación diaria de redes"),
             ("Sergio Barranco", "Líder de diseño · recibe y asigna las solicitudes de diseño"),
             ("Kevin Moss", "Foto y video de activaciones y eventos")]
    r = section(ws, r, "Quién es quién en social", 5)
    mini_table(ws, r, ["Persona", "Cargo", "Rol en social"],
               [(n, (equipo.get(n) or {}).get("Cargo"), rol) for n, rol in voces], wrap_cols=(2,))
    for col, w in zip("ABCDE", (30, 44, 60, 14, 4)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].width = 44
    return ws


def sheet_linea(wb):
    ws = wb.create_sheet("03_Línea_Editorial")
    layout.banner(ws, "Línea editorial", 6)
    r = 5
    r = section(ws, r, "Cadencia", 6)
    r = mini_table(ws, r, ["Contenido", "Formato", "Frecuencia", "Día", "Plataforma", "Aprobación"], [
        ("Repost del día", "Historia", "Diaria", "Todos", "Instagram", "—"),
        ("Post semanal", "Post / carrusel / reel", "1 por semana", "Martes", "Instagram", "—"),
        ("Post semanal 2", "Post / carrusel / reel", "Opcional", "Jueves", "Instagram", "—"),
        ("Historia larga", "Historia larga", "1 por semana", "Viernes", "Instagram", "—"),
        ("Cruce con LinkedIn", "Post / documento", "Mínimo 1 al mes", "Primer martes del mes", "Instagram + LinkedIn",
         "PR · Jocelyn Santiago"),
        ("Libre", "Cualquiera", "Hasta 4 publicaciones por día", "—", "—", "Si cruza a LinkedIn, PR"),
    ])
    r = section(ws, r, "Tipos de contenido", 6)
    r = mini_table(ws, r, ["Tipo", "Qué entra", "Formatos", "¿Cruza a LinkedIn?", "Insumo"], [
        ("Repost", "Contenido de creadores, clientes o medios que menciona a MCR", "Historia", "No", "Menciones y etiquetas"),
        ("Trabajo con marcas", "Campañas y resultados con clientes", "Reel, carrusel", "Sí, con resultados",
         "Portafolio y campañas"),
        ("Activaciones en vivo", "Eventos presenciales, OOH, experiencias", "Historia, reel", "Sí, el resumen",
         "Foto y video de Kevin"),
        ("Eventos de creadores", "Encuentros, viajes y eventos del roster", "Historia, reel, carrusel", "Opcional",
         "Talent Management"),
        ("Premios a la agencia", "Premios, finales, rankings y membresías", "Post, carrusel", "Sí", "Prensa y premios"),
        ("Artículos de Sergi", "Columnas, entrevistas y notas", "Carrusel, artículo LinkedIn", "Sí", "Prensa"),
        ("Pláticas de Sergi", "Paneles, conferencias y podcasts", "Reel, carrusel", "Sí", "Agenda de Sergi"),
        ("Otro", "Lo que no está previsto", "Cualquiera", "Según el caso", "—"),
    ], wrap_cols=(1, 2))
    r = section(ws, r, "Formatos y medidas", 6)
    r = mini_table(ws, r, ["Formato", "Medida"], [
        ("Historia / reel", "1080×1920 · 9:16"), ("Post / carrusel", "1080×1350 · 4:5 (hasta 10 láminas)"),
        ("Post cuadrado", "1080×1080 · 1:1"), ("LinkedIn imagen", "1200×627"),
        ("LinkedIn documento", "PDF 1080×1350 · 4:5"), ("Video horizontal", "1920×1080 · 16:9")])
    r = section(ws, r, "Flujo", 6)
    mini_table(ws, r, ["Paso", "Quién", "Dónde"], [
        ("1. Solicitud de diseño", "Sergi, Jocelyn o Dulce", "10_Solicitudes_Diseño"),
        ("2. Asignación", "Sergio Barranco (líder de diseño)", "10_Solicitudes_Diseño · Asignado a"),
        ("3. Diseño", "Sergio, Elena, Diego Galaviz o diseñador 3", "10_Solicitudes_Diseño · Status"),
        ("4. Revisión", "Quien solicitó", "10_Solicitudes_Diseño · Notas / correcciones"),
        ("5. Aprobación PR (si cruza a LinkedIn)", "Jocelyn Santiago", "Calendario · Aprobación PR"),
        ("6. Programación y publicación", "Dulce Valentín", "Calendario · Estado"),
    ])
    for col, w in zip("ABCDEF", (30, 46, 28, 22, 24, 22)):
        ws.column_dimensions[col].width = w


CAL_COLS = [("fecha", "Fecha", 12), ("dia", "Día", 6), ("semana", "Semana", 8), ("slot", "Slot", 6),
            ("plataforma", "Plataforma", 18), ("formato", "Formato", 16), ("tipo", "Tipo de contenido", 22),
            ("pieza", "Pieza / tema", 34), ("copy", "Copy", 44), ("linkedin", "Cruce LinkedIn", 10),
            ("aprobacion", "Aprobación PR", 16), ("solicitud", "ID solicitud", 12), ("estado_diseno", "Estado diseño", 14),
            ("responsable", "Responsable", 18), ("estado", "Estado", 14), ("asset", "Link asset", 24),
            ("link", "Link publicación", 24), ("notas", "Notas", 30)]


def month_plan(year, month):
    """Slots fijos: historia repost diaria; post martes; post 2 opcional jueves; historia larga viernes;
    el primer post del mes cruza a LinkedIn con aprobación de PR."""
    rows = []
    days = calendar.monthrange(year, month)[1]
    first_tuesday = next(d for d in range(1, 8) if dt.date(year, month, d).weekday() == 1)
    for d in range(1, days + 1):
        date = dt.date(year, month, d)
        wd = date.weekday()
        planned = [dict(plataforma="Instagram", formato="Historia", tipo="Repost", pieza="Repost del día",
                        responsable="Dulce Valentín", estado="Por definir")]
        if wd == 1:
            li = d == first_tuesday
            planned.append(dict(plataforma="Instagram + LinkedIn" if li else "Instagram", formato="Post",
                                pieza="Post semanal" + (" · cruce LinkedIn" if li else ""), linkedin="Sí" if li else "No",
                                aprobacion="Pendiente" if li else None, estado="Por definir"))
        elif wd == 3:
            planned.append(dict(plataforma="Instagram", formato="Post", pieza="Post semanal 2 (opcional)",
                                linkedin="No", estado="Por definir"))
        elif wd == 4:
            planned.append(dict(plataforma="Instagram", formato="Historia larga", pieza="Historia larga semanal",
                                estado="Por definir"))
        for s in range(1, SLOTS + 1):
            base = dict(fecha=date, dia=DIAS[wd], semana=date.isocalendar()[1], slot=s)
            if s <= len(planned):
                base.update(planned[s - 1])
            rows.append(base)
    return rows


def sheet_month(wb, name, year, month, req_sheet):
    ws = wb.create_sheet(name)
    layout.banner(ws, f"{MESES[month - 1]} {year}", len(CAL_COLS))
    hdr = 4
    L = {f: get_column_letter(i) for i, (f, _, _) in enumerate(CAL_COLS, 1)}
    for i, (f, label, w) in enumerate(CAL_COLS, 1):
        cell(ws, hdr, i, label, bold=True, color=WHITE, fill=PURPLE, wrap=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    rows = month_plan(year, month)
    for k, row in enumerate(rows):
        r = hdr + 1 + k
        for i, (f, _, _) in enumerate(CAL_COLS, 1):
            v = row.get(f)
            if f == "estado_diseno":
                sid = f"{L['solicitud']}{r}"
                v = (f"=IF({sid}=\"\",\"\",IFERROR(INDEX('{req_sheet}'!$S$5:$S$1000,"
                     f"MATCH({sid},'{req_sheet}'!$A$5:$A$1000,0)),\"\"))")
            x = cell(ws, r, i, v, color=MUTED if f in ("dia", "semana", "slot") and row["slot"] > 1 else BLACK,
                     align="center" if f in ("dia", "semana", "slot", "linkedin") else None)
            if f == "fecha":
                x.number_format = "DD/MM/YYYY"
                if row["slot"] > 1:
                    x.font = Font(name=FONT, size=10, color=MUTED)
            if f == "estado_diseno":
                x.font = Font(name=FONT, size=10, color=PURPLE_600)
            x.border = Border(top=DAY_TOP if row["slot"] == 1 else THIN, bottom=THIN)
    last = hdr + len(rows)
    t = Table(displayName=f"Cal_{year}_{month:02d}", ref=f"A{hdr}:{get_column_letter(len(CAL_COLS))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(hdr + 1, 5)
    ws.row_dimensions[hdr].height = 30
    rng = lambda f: f"{L[f]}{hdr + 1}:{L[f]}{last}"
    full = f"A{hdr + 1}:{get_column_letter(len(CAL_COLS))}{last}"
    # día par / impar y fin de semana
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'OR($B{hdr + 1}="Sáb",$B{hdr + 1}="Dom")'],
                                                    fill=PatternFill("solid", fgColor="EDE3FF")))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f"MOD(DAY($A{hdr + 1}),2)=0"],
                                                    fill=PatternFill("solid", fgColor="FAF7FF")))
    # LinkedIn sin aprobación de PR
    ws.conditional_formatting.add(rng("aprobacion"), FormulaRule(
        formula=[f'AND(${L["linkedin"]}{hdr + 1}="Sí",${L["aprobacion"]}{hdr + 1}<>"Aprobado")'],
        fill=PatternFill("solid", fgColor="FFD9D9"), font=Font(name=FONT, bold=True, color="B00020")))
    ws.conditional_formatting.add(rng("estado"), CellIsRule(operator="equal", formula=['"Publicado"'],
                                                            fill=PatternFill("solid", fgColor="D6F8E9"),
                                                            font=Font(name=FONT, bold=True, color="00794A")))
    ws.conditional_formatting.add(rng("estado"), CellIsRule(operator="equal", formula=['"Programado"'],
                                                            font=Font(name=FONT, bold=True, color=PURPLE)))
    for f, opts in (("plataforma", PLATAFORMAS), ("formato", FORMATOS), ("tipo", TIPOS), ("linkedin", ["Sí", "No"]),
                    ("aprobacion", APROBACION), ("responsable", RESPONSABLES), ("estado", ESTADOS)):
        dv_list(ws, rng(f), opts)
    dv = DataValidation(type="list", formula1=f"'{req_sheet}'!$A$5:$A$1000", allow_blank=True)
    dv.add(rng("solicitud"))
    ws.add_data_validation(dv)
    return len(rows)


REQ_COLS = [("id", "ID", 9), ("pieza", "Nombre de la pieza", 30), ("tipo", "Tipo de contenido", 20),
            ("plataforma", "Plataforma", 16), ("formato", "Formato", 14), ("solicita", "Quién solicita", 18),
            ("para", "A quién solicita", 18), ("asignado", "Asignado a", 18), ("f_solicitud", "Fecha de solicitud", 13),
            ("f_publicacion", "Fecha de publicación", 13), ("f_entrega", "Fecha de entrega", 13),
            ("prioridad", "Prioridad", 11), ("descripcion", "Descripción", 46), ("copy", "Copy en imagen", 36),
            ("referencia", "Referencia", 28), ("medidas", "Medidas", 26), ("linkedin", "Cruce LinkedIn", 10),
            ("aprobacion", "Aprobación PR", 16), ("status", "Status", 13), ("asset", "Link del asset", 26),
            ("notas", "Notas / correcciones", 34), ("final", "Entrega final", 26)]


def sheet_requests(wb, name, n=300):
    ws = wb.create_sheet(name)
    layout.banner(ws, "Solicitudes de diseño · social", len(REQ_COLS))
    hdr = 4
    for i, (f, label, w) in enumerate(REQ_COLS, 1):
        cell(ws, hdr, i, label, bold=True, color=WHITE, fill=PURPLE, wrap=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    L = {f: get_column_letter(i) for i, (f, _, _) in enumerate(REQ_COLS, 1)}
    assert L["status"] == "S", "el calendario busca el status en la columna S"
    for k in range(n):
        r = hdr + 1 + k
        for i, (f, _, _) in enumerate(REQ_COLS, 1):
            v = f"DS-{k + 1:03d}" if f == "id" else LIDER_DISENO if f == "para" else None
            x = cell(ws, r, i, v, color=MUTED if f == "id" else BLACK, wrap=f in ("descripcion", "copy", "notas"))
            x.border = Border(bottom=THIN)
            if f.startswith("f_"):
                x.number_format = "DD/MM/YYYY"
    last = hdr + n
    t = Table(displayName="Solicitudes_Diseno", ref=f"A{hdr}:{get_column_letter(len(REQ_COLS))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(hdr + 1, 3)
    ws.row_dimensions[hdr].height = 30
    rng = lambda f: f"{L[f]}{hdr + 1}:{L[f]}{last}"
    for f, opts in (("tipo", TIPOS), ("plataforma", PLATAFORMAS), ("formato", FORMATOS), ("solicita", SOLICITANTES),
                    ("para", [LIDER_DISENO]), ("asignado", DISENO), ("prioridad", PRIORIDAD), ("medidas", MEDIDAS),
                    ("linkedin", ["Sí", "No"]), ("aprobacion", APROBACION), ("status", STATUS_DISENO)):
        dv_list(ws, rng(f), opts)
    colors = {"Urgente": ("FFD9D9", "B00020"), "Alta": ("FFE8CC", "8A4B00"), "Terminada": ("D6F8E9", "00794A"),
              "Aprobado": ("D6F8E9", "00794A"), "Bloqueo": ("FFD9D9", "B00020"), "Cambios": ("FFF3C4", "7A5A00")}
    for f in ("prioridad", "status"):
        for val, (bg, fg) in colors.items():
            ws.conditional_formatting.add(rng(f), CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                             fill=PatternFill("solid", fgColor=bg),
                                                             font=Font(name=FONT, bold=True, color=fg)))
    ws.conditional_formatting.add(rng("aprobacion"), FormulaRule(
        formula=[f'AND(${L["linkedin"]}{hdr + 1}="Sí",${L["aprobacion"]}{hdr + 1}<>"Aprobado")'],
        fill=PatternFill("solid", fgColor="FFD9D9"), font=Font(name=FONT, bold=True, color="B00020")))
    return ws


def sheet_summary(ws, months, req_sheet):
    layout.banner(ws, "Resumen", 8)
    hdr = 5
    heads = ["Mes", "Publicaciones planeadas", "Programadas", "Publicadas", "Cruces LinkedIn", "PR pendiente",
             "Slots libres"]
    for i, h in enumerate(heads, 1):
        cell(ws, hdr, i, h, bold=True, color=WHITE, fill=PURPLE, wrap=True)
    for k, (name, label) in enumerate(months):
        r = hdr + 1 + k
        s = f"'{name}'"
        cell(ws, r, 1, label, bold=True, color=PURPLE)
        ws.cell(r, 1).hyperlink = f"#{s}!A1"
        formulas = [f'=COUNTIF({s}!$H$5:$H$200,"?*")', f'=COUNTIF({s}!$O$5:$O$200,"Programado")',
                    f'=COUNTIF({s}!$O$5:$O$200,"Publicado")', f'=COUNTIF({s}!$J$5:$J$200,"Sí")',
                    f'=COUNTIFS({s}!$J$5:$J$200,"Sí",{s}!$K$5:$K$200,"<>Aprobado")',
                    f'=COUNTIFS({s}!$A$5:$A$200,"<>",{s}!$H$5:$H$200,"")']
        for i, f in enumerate(formulas, 2):
            x = cell(ws, r, i, f, bold=True, color=PURPLE, size=12, align="right")
            x.border = Border(bottom=THIN)
    r = hdr + len(months) + 3
    q = f"'{req_sheet}'"
    cell(ws, r, 1, "SOLICITUDES DE DISEÑO", bold=True, color=WHITE, fill=BLACK)
    ws.cell(r, 2).fill = PatternFill("solid", fgColor=BLACK)
    r += 1
    for st in STATUS_DISENO:
        cell(ws, r, 1, st)
        x = cell(ws, r, 2, f'=COUNTIF({q}!$S$5:$S$1000,"{st}")', bold=True, color=PURPLE, size=12, align="right")
        x.border = Border(bottom=THIN)
        r += 1
    cell(ws, r, 1, "Urgentes abiertas")
    cell(ws, r, 2, f'=COUNTIFS({q}!$L$5:$L$1000,"Urgente",{q}!$S$5:$S$1000,"<>Terminada")', bold=True, color=PURPLE,
         size=12, align="right")
    for col, w in zip("ABCDEFG", (24, 16, 14, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[hdr].height = 32


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--maestro", required=True)
    ap.add_argument("--inicio", default="2026-11-01")
    ap.add_argument("--meses", type=int, default=6)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    co = company(args.maestro)
    start = dt.date.fromisoformat(args.inicio)
    wb = layout.new_workbook()
    idx = wb.create_sheet("00_Índice")
    res = wb.create_sheet("01_Resumen")
    sheet_marca(wb, co)
    sheet_linea(wb)
    req_sheet = f"{4 + args.meses:02d}_Solicitudes_Diseño"
    months, counts = [], []
    y, m = start.year, start.month
    for i in range(args.meses):
        name = f"{4 + i:02d}_{MESES[m - 1][:3]}_{y}"
        n = sheet_month(wb, name, y, m, req_sheet)
        months.append((name, f"{MESES[m - 1]} {y}"))
        counts.append(n)
        m = 1 if m == 12 else m + 1
        y = y + 1 if m == 1 else y
    sheet_requests(wb, req_sheet)
    sheet_summary(res, months, req_sheet)
    entries = [("01_Resumen", "Avance por mes y estado de diseño", None), ("02_Marca", "Identidad, paleta, cifras, servicios, casos y premios", None),
               ("03_Línea_Editorial", "Cadencia, tipos de contenido, medidas y flujo", None)]
    entries += [(name, f"Calendario · {label}", None) for name, label in months]
    entries += [(req_sheet, "Solicitudes de diseño para social", None)]
    layout.index_sheet(idx, entries, dt.date.today().strftime("%d/%m/%Y"))
    idx["A5"] = "Biblia de social"
    wb.save(args.out)
    print(f"OK · {args.out}")


if __name__ == "__main__":
    main()
