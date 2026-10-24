"""Capa visual del maestro: paleta MCR (tema del sitio), tablas, índice y resumen."""
import re

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.writer.theme import theme_xml

# Tokens de assets/global.css del tema
PURPLE = "6700F2"
PURPLE_600 = "5500CC"
PURPLE_400 = "8833FF"
PURPLE_200 = "C7A8FF"
PURPLE_50 = "F3ECFF"
BLACK = "0F0F0F"
SURFACE = "1E1E1E"
MUTED = "8A8A8A"
GREEN = "00DF82"
WHITE = "FFFFFF"
FONT = "Arial"

TIER_STYLE = {"AAA": (PURPLE, WHITE, True), "Mid": (PURPLE_200, BLACK, True), "Rising": (PURPLE_50, PURPLE_600, False)}


def new_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    theme = theme_xml.replace('<a:srgbClr val="4F81BD"/>', f'<a:srgbClr val="{PURPLE}"/>', 1)
    theme = theme.replace('<a:srgbClr val="1F497D"/>', f'<a:srgbClr val="{BLACK}"/>', 1)
    theme = theme.replace('<a:srgbClr val="0000FF"/>', f'<a:srgbClr val="{PURPLE}"/>', 1)
    theme = re.sub(r'<a:latin typeface="(Cambria|Calibri)"', f'<a:latin typeface="{FONT}"', theme)
    wb.loaded_theme = theme.encode("utf-8")
    return wb


def banner(ws, title, width_cols):
    last = get_column_letter(max(width_cols, 4))
    for col in range(1, max(width_cols, 4) + 1):
        ws.cell(1, col).fill = PatternFill("solid", fgColor=BLACK)
        ws.cell(2, col).fill = PatternFill("solid", fgColor=PURPLE)
    ws["A1"] = title.upper()
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 4
    ws.row_dimensions[3].height = 10
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PURPLE
    return last


def table_sheet(wb, name, title, columns, rows, widths=None, link_cols=(), link_text=None, int_cols=(), pct_cols=(),
                wrap_cols=(), center_cols=(), freeze_cols=2, empty_rows=0, date_cols=(), lists=None):
    """columns: [(campo, etiqueta)]. Encabezado en la fila 4 y datos desde la 5."""
    labels = [label for _, label in columns]
    dupes = {l for l in labels if labels.count(l) > 1}
    if dupes:
        raise ValueError(f"{name}: encabezados repetidos {dupes}")
    ws = wb.create_sheet(name)
    banner(ws, title, len(columns))
    hdr = 4
    body = Font(name=FONT, size=10, color=BLACK)
    link_font = Font(name=FONT, size=10, color=PURPLE, underline="single")
    line = Border(bottom=Side(style="thin", color="E4DCF5"))
    for i, (_, label) in enumerate(columns, 1):
        c = ws.cell(hdr, i, label)
        c.font = Font(name=FONT, bold=True, size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor=PURPLE)
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=0)
    for r_i, row in enumerate(rows, hdr + 1):
        for c_i, (field, _) in enumerate(columns, 1):
            v = row.get(field)
            if isinstance(v, (list, tuple)):
                v = ", ".join(map(str, v))
            cell = ws.cell(r_i, c_i, v)
            cell.font = body
            cell.border = line
            cell.alignment = Alignment(vertical="top", wrap_text=field in wrap_cols,
                                       horizontal="center" if field in center_cols else None)
            if field in link_cols and isinstance(v, str) and re.match(r"(https?://|www\.)", v):
                first = v.split(",")[0].strip()
                cell.hyperlink = first
                cell.value = link_text or re.sub(r"^https?://(www\.|m\.)?", "", first).rstrip("/")
                cell.font = link_font
            if field in int_cols and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
            if field in pct_cols and isinstance(v, (int, float)):
                cell.number_format = "0.00"
    for r_i in range(hdr + len(rows) + 1, hdr + len(rows) + empty_rows + 1):
        for c_i in range(1, len(columns) + 1):
            ws.cell(r_i, c_i).font = body
            ws.cell(r_i, c_i).border = line
    for r_i in range(hdr + 1, hdr + len(rows) + empty_rows + 1):
        for c_i, (field, _) in enumerate(columns, 1):
            if field in date_cols:
                ws.cell(r_i, c_i).number_format = "DD/MM/YYYY"
    last_row = hdr + max(len(rows) + empty_rows, 1)
    last_col = get_column_letter(len(columns))
    if rows or empty_rows:
        t = Table(displayName=f"T_{re.sub(r'[^A-Za-z0-9]', '_', name)}", ref=f"A{hdr}:{last_col}{last_row}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = ws.cell(hdr + 1, freeze_cols + 1 if len(columns) > 5 else 1)
    widths = widths or {}
    for i, (field, label) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(field, max(11, min(34, len(label) + 5)))
    ws.row_dimensions[hdr].height = 30
    for field, options in (lists or {}).items():
        col = get_column_letter([f for f, _ in columns].index(field) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{col}{hdr + 1}:{col}{last_row + 500}")
        ws.add_data_validation(dv)
    for i, (field, _) in enumerate(columns, 1):
        if field == "tier":
            col = get_column_letter(i)
            for tier, (bg, fg, bold) in TIER_STYLE.items():
                ws.conditional_formatting.add(
                    f"{col}{hdr + 1}:{col}{last_row + 500}",
                    CellIsRule(operator="equal", formula=[f'"{tier}"'], fill=PatternFill("solid", fgColor=bg),
                               font=Font(name=FONT, color=fg, bold=bold)))
    return ws


def index_sheet(ws, entries, date_text):
    banner(ws, "MCR Agency", 4)
    ws["A1"].font = Font(name=FONT, bold=True, size=22, color=WHITE)
    ws.row_dimensions[1].height = 48
    ws["A5"] = "Base maestra"
    ws["A5"].font = Font(name=FONT, bold=True, size=14, color=PURPLE)
    ws["A6"] = f"Actualizada al {date_text}"
    ws["A6"].font = Font(name=FONT, size=10, color=MUTED)
    r = 8
    for col, label in enumerate(("Hoja", "Contenido", "Registros"), 1):
        c = ws.cell(r, col, label)
        c.font = Font(name=FONT, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=PURPLE)
    for name, desc, n in entries:
        r += 1
        c = ws.cell(r, 1, name)
        c.hyperlink = f"#'{name}'!A1"
        c.font = Font(name=FONT, bold=True, color=PURPLE)
        ws.cell(r, 2, desc).font = Font(name=FONT, color=BLACK)
        nc = ws.cell(r, 3, n)
        nc.font = Font(name=FONT, color=MUTED)
        nc.number_format = "#,##0"
        for col in (1, 2, 3):
            ws.cell(r, col).border = Border(bottom=Side(style="thin", color="E4DCF5"))
            if r % 2 == 0:
                ws.cell(r, col).fill = PatternFill("solid", fgColor=PURPLE_50)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 4


def summary_sheet(ws, blocks):
    banner(ws, "Resumen", 6)
    r = 5
    col_pairs = [(1, 2), (4, 5)]
    pos = {0: 5, 1: 5}
    for i, (title, items) in enumerate(blocks):
        side = 0 if i % 2 == 0 else 1
        c1, c2 = col_pairs[side]
        r = pos[side]
        for col in (c1, c2):
            ws.cell(r, col).fill = PatternFill("solid", fgColor=BLACK)
        ws.cell(r, c1, title.upper()).font = Font(name=FONT, bold=True, color=WHITE, size=10)
        r += 1
        for label, formula in items:
            a = ws.cell(r, c1, label)
            a.font = Font(name=FONT, color=BLACK, size=10)
            b = ws.cell(r, c2, formula)
            b.number_format = "#,##0"
            b.font = Font(name=FONT, bold=True, color=PURPLE, size=12)
            b.alignment = Alignment(horizontal="right")
            for col in (c1, c2):
                ws.cell(r, col).border = Border(bottom=Side(style="thin", color="E4DCF5"))
            r += 1
        pos[side] = r + 1
    for col, w in zip("ABCDE", (40, 16, 4, 40, 16)):
        ws.column_dimensions[col].width = w
